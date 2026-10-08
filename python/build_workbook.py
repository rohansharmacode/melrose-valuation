"""Build model/Melrose_DCF.xlsx with live formulas.

Every forecast and valuation cell is an Excel formula that traces back to
the Inputs and Scenarios sheets, so changing an assumption re-prices the
whole model. python/dcf.py is an independent replica used to check it.
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = Path(__file__).resolve().parents[1] / "model" / "Melrose_DCF.xlsx"

FONT = "Arial"
BLUE, BLACK, GREEN, GREY = "0000FF", "000000", "008000", "666666"
YELLOW = PatternFill("solid", fgColor="FFFF00")
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
SECTION_FILL = PatternFill("solid", fgColor="D9E1F2")
THIN = Side(style="thin", color="999999")

GBP = '#,##0;(#,##0);"-"'
GBP1 = '#,##0.0;(#,##0.0);"-"'
PCT = '0.0%;(0.0%);"-"'
PCT2 = '0.00%;(0.00%);"-"'
MULT = '0.0"x"'
PENCE = '#,##0"p"'
PENCE1 = '#,##0.0"p"'
FACTOR = "0.000"

YEARS = list(range(2025, 2036))           # C = 2025A ... M = 2035E
COLS = [get_column_letter(3 + i) for i in range(len(YEARS))]
COL = dict(zip(YEARS, COLS))


def font(color=BLACK, bold=False, size=10, italic=False):
    return Font(name=FONT, color=color, bold=bold, size=size, italic=italic)


def put(ws, ref, value, fmt=None, color=BLACK, bold=False, fill=None,
        italic=False, align=None):
    cell = ws[ref]
    cell.value = value
    cell.font = font(color, bold, italic=italic)
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = fill
    if align:
        cell.alignment = Alignment(horizontal=align)
    return cell


def title(ws, text, sub=None):
    ws["A1"].value = text
    ws["A1"].font = font(bold=True, size=14)
    if sub:
        ws["A2"].value = sub
        ws["A2"].font = font(GREY, italic=True)


def section(ws, row, text, last_col="E"):
    for col in range(1, ws[f"{last_col}1"].column + 1):
        ws.cell(row=row, column=col).fill = SECTION_FILL
    put(ws, f"A{row}", text, bold=True)


def header_row(ws, row, labels, start_col=1):
    for i, label in enumerate(labels):
        cell = ws.cell(row=row, column=start_col + i, value=label)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="center" if i else "left",
                                   wrap_text=True, vertical="center")


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


# --------------------------------------------------------------- Inputs
INPUT_ROWS = {}


def build_inputs(wb):
    ws = wb.create_sheet("Inputs")
    title(ws, "Inputs and sources",
          "Blue = typed-in data with its source. Black = formula. "
          "£m unless stated.")
    header_row(ws, 3, ["Item", "Value", "Unit", "Source / reasoning"])
    rows = [
        ("Market data", None),
        ("price", "Share price", 450.0, "p", PENCE1,
         "LSE closing price, 7 Oct 2026 (lse.co.uk)"),
        ("shares", "Shares in issue", 1246, "m", GBP,
         "H1 2026 results (31 Jul 2026): excludes treasury shares, 30 Jun 2026"),
        ("mcap", "Market capitalisation", "={price}*{shares}/100", "£m", GBP, "Formula"),
        ("net_debt", "Net debt, 30 Jun 2026", 1530, "£m", GBP,
         "H1 2026 results; excludes lease liabilities"),
        ("leases", "Lease liabilities", 330, "£m", GBP,
         "FY2025 results, 31 Dec 2025 (£31m current + £299m non-current)"),
        ("pension", "Net pension deficit, 30 Jun 2026", 11, "£m", GBP,
         "H1 2026 results"),
        ("ev", "Enterprise value incl. leases", "={mcap}+{net_debt}+{leases}+{pension}", "£m", GBP,
         "Formula: pairs with EBITDA, which adds back lease depreciation"),
        ("low52", "52-week low", 409.9, "p", PENCE1,
         "31 Jul 2026, H1 results day (lse.co.uk)"),
        ("high52", "52-week high", 682.6, "p", PENCE1,
         "19 Feb 2026, before FY2025 results (lse.co.uk)"),
        ("eps26", "Consensus adjusted EPS 2026E", 35, "p", PENCE1,
         "S&P Global consensus via stockanalysis.com, 7 Oct 2026, 14-15 analysts"),
        ("eps27", "Consensus adjusted EPS 2027E", 45, "p", PENCE1, "As above"),
        ("tgt_avg", "Analyst price target, average", 663, "p", PENCE1,
         "As above (15 analysts: 10 buy, 3 hold, 2 sell)"),
        ("tgt_low", "Analyst price target, low", 450, "p", PENCE1, "As above"),
        ("tgt_high", "Analyst price target, high", 825, "p", PENCE1, "As above"),
        ("FY2026 guidance mid-points (excludes Garden Grove; at US$1.37/£)",
         None),
        ("g_eng_rev", "Engines revenue", 1750, "£m", GBP,
         "FY2025 results (27 Feb 2026): £1,700-1,800m; reaffirmed 31 Jul 2026"),
        ("g_eng_op", "Engines adjusted operating profit", 580, "£m", GBP,
         "Guidance £565-595m"),
        ("g_af_rev", "Airframes revenue", 2100, "£m", GBP,
         "Guidance £2,050-2,150m"),
        ("g_af_op", "Airframes adjusted operating profit", 180, "£m", GBP,
         "Guidance £170-190m"),
        ("g_op", "Group adjusted operating profit", 725, "£m", GBP,
         "Guidance £700-750m"),
        ("g_vc", "Variable consideration (non-cash RRSP profit)", 360, "£m",
         GBP, "Guidance £340-380m, depending on engine build rates"),
        ("g_fcf", "Free cash flow after interest and tax", 175, "£m", GBP,
         "Guidance £150-200m, including a £30-50m factoring inflow"),
        ("g_corp", "Central costs (implied)", "={g_eng_op}+{g_af_op}-{g_op}", "£m", GBP,
         "Formula: divisions less group"),
        ("2029 targets (at US$1.25/£)", None),
        ("t_rev", "Revenue", 5000, "£m", GBP, "FY2025 results"),
        ("t_op", "Adjusted operating profit", 1200, "£m", GBP,
         "FY2025 results: '£1.2+ billion'"),
        ("t_fcf", "Free cash flow after interest and tax", 600, "£m", GBP,
         "FY2025 results; reaffirmed in H1 2026 results"),
        ("fx_sens", "Profit change per 1 cent move in US$", 0.007, "%", PCT,
         "H1 2026 results: 1 cent = £5m (0.7%) of adjusted operating profit"),
        ("fx_guid", "Exchange rate in 2026 guidance", 1.37, "US$/£", "0.00",
         "FY2025 and H1 2026 results"),
        ("fx_tgt", "Exchange rate in 2029 targets", 1.25, "US$/£", "0.00",
         "FY2025 results"),
        ("fx_factor", "Restates targets at the guidance rate",
         "=1-{fx_sens}*({fx_guid}-{fx_tgt})*100", "x", FACTOR,
         "Formula: a weaker dollar shrinks sterling profit"),
        ("t_fcf_137", "2029 FCF target at US$1.37", "={t_fcf}*{fx_factor}", "£m", GBP,
         "Formula: the like-for-like yardstick for this model"),
        ("Garden Grove (US transparencies plant, chemical tank incident May 2026)",
         None),
        ("gg_rev", "Lost revenue in 2026", 46, "£m", GBP,
         "H1 actual £16m + H2 estimate £30m: about £6m a month at half "
         "capacity; full restart slipped past the 28 Sep target"),
        ("gg_op", "Lost adjusted operating profit in 2026", 39, "£m", GBP,
         "H1 actual £9m + H2 estimate £30m"),
        ("H1 2026 actuals (the first half is already in the 30 Jun net debt)",
         None),
        ("h1_op", "Adjusted operating profit", 347, "£m", GBP,
         "H1 2026 results"),
        ("h1_vc", "Variable consideration", 206, "£m", GBP, "H1 2026 results"),
        ("h1_fcf", "Free cash flow before factoring", 28, "£m", GBP,
         "H1 2026 results"),
        ("h1_it", "Net interest and tax paid", 71, "£m", GBP,
         "H1 2026 results"),
        ("h1_gg", "Garden Grove cash costs", 5, "£m", GBP, "H1 2026 results"),
        ("h1_fcff", "First-half unlevered FCF (model basis)",
         "={h1_fcf}+{h1_it}+{h1_gg}-{tax}*({h1_op}-{h1_vc})", "£m", GBP,
         "Formula: adds back interest, tax and Garden Grove; taxes as below"),
        ("Forecast assumptions (all scenarios)", None),
        ("da_pct", "Depreciation and amortisation", 0.038, "% revenue", PCT,
         "FY2025: £138m on £3,589m revenue"),
        ("capex_pct", "Capital expenditure", 0.030, "% revenue", PCT,
         "FY2025: 2.6%; H1 2026: 2.8%; 3.0% allows for growth"),
        ("lease_pct", "Lease payments", 0.009, "% revenue", PCT,
         "FY2025: £31m"),
        ("wc_pct", "Working capital, excluding VC", 0.10,
         "% of revenue increase", PCT, "Assumption: stock builds with output"),
        ("corp_g", "Central cost growth", 0.03, "% a year", PCT,
         "Assumption: inflation"),
        ("pen_cash", "Pension cash contributions", 5, "£m a year", GBP,
         "H1 2026: £2m (FY2025: £22m); deficit now only £11m"),
        ("fade_to", "Growth that 2030-35 fades towards", 0.025, "% a year",
         PCT, "Assumption: long-run nominal growth"),
        ("tax", "Tax rate", 0.204, "%", PCT,
         "FY2025 adjusted effective rate; applied to profit after "
         "non-cash VC"),
        ("gtf26", "GTF powder-metal cash cost 2026", 50, "£m", GBP,
         "Guided c.£50m"),
        ("gtf27", "GTF powder-metal cash cost 2027", 30, "£m", GBP,
         "Total programme within c.£200m (2024 £35m, 2025 £68m)"),
        ("gtf28", "GTF powder-metal cash cost 2028", 15, "£m", GBP,
         "GTF programmes turn cash positive in 2028 (H1 2026 results)"),
        ("ro26", "Restructuring and other, 2026", 25, "£m", GBP,
         "H1 2026 restructuring £6m; non-cash contract releases ~£10m a year"),
        ("ro27", "Restructuring and other, 2027", 15, "£m", GBP, "Tapering"),
        ("ro28", "Restructuring and other, 2028", 5, "£m", GBP, "Tapering"),
        ("Discount rate", None),
        ("rf", "Risk-free rate (UK 10-year gilt)", 0.0548, "%", PCT2,
         "Trading Economics, 8 Oct 2026"),
        ("erp", "Equity risk premium", 0.05, "%", PCT2, "Assumption"),
        ("beta", "Beta", 1.0, "x", "0.00",
         "5-year beta 0.89 (stockanalysis.com); 1.0 for aerospace "
         "cyclicality"),
        ("ke", "Cost of equity", "={rf}+{beta}*{erp}", "%", PCT2, "Formula: CAPM"),
        ("spread", "Credit spread over gilts", 0.015, "%", PCT2,
         "Assumption for a BBB-type borrower"),
        ("tax_int", "Tax relief on interest", 0.25, "%", PCT,
         "UK main corporation tax rate"),
        ("kd", "After-tax cost of debt", "=({rf}+{spread})*(1-{tax_int})", "%", PCT2,
         "Formula"),
        ("we", "Equity weight", "={mcap}/({mcap}+{net_debt})", "%", PCT, "Formula: market values"),
        ("wacc", "WACC", "={we}*{ke}+(1-{we})*{kd}", "%", PCT2, "Formula"),
        ("g", "Terminal growth", 0.025, "%", PCT2,
         "Assumption: below long-run UK/US nominal GDP growth"),
    ]
    # pass 1: give every item a row, so formulas can name items by key
    r = 4
    placed = []
    for item in rows:
        placed.append((r, item))
        if item[1] is not None:
            INPUT_ROWS[item[0]] = r
        r += 1
    refs = {k: f"B{v}" for k, v in INPUT_ROWS.items()}
    # pass 2: write
    for r, item in placed:
        if item[1] is None:
            section(ws, r, item[0], "D")
            continue
        key, label, value, unit, fmt, source = item
        is_formula = isinstance(value, str) and value.startswith("=")
        if is_formula:
            value = value.format(**refs)
        put(ws, f"A{r}", label)
        put(ws, f"B{r}", value, fmt, BLACK if is_formula else BLUE)
        put(ws, f"C{r}", unit, color=GREY)
        put(ws, f"D{r}", source, color=GREY)
    widths(ws, {"A": 46, "B": 11, "C": 18, "D": 92})
    ws.freeze_panes = "A4"
    return ws


def I(key, absolute=True):
    row = INPUT_ROWS[key]
    return f"Inputs!$B${row}" if absolute else f"Inputs!B{row}"


# ------------------------------------------------------------ Scenarios
SCEN_COL = {"Bear": "B", "Base": "C", "Bull": "D"}
SCEN_ROWS = {}


def build_scenarios(wb):
    ws = wb.create_sheet("Scenarios")
    title(ws, "Scenario levers",
          "Yellow = the judgement calls that drive the valuation.")
    header_row(ws, 3, ["Driver", "Bear", "Base", "Bull", "Why"])
    levers = [
        ("eng_g", "Engines revenue growth, 2027-29", (0.05, 0.08, 0.10), PCT,
         "Management: high single-digit growth a year to 2029"),
        ("eng_m29", "Engines margin reached by 2029", (0.32, 0.35, 0.37), PCT,
         "Management: mid-to-high 30s%. 2025 actual 31.9%; 2026 guidance 33.1%"),
        ("af_g", "Airframes revenue growth, 2027-29", (0.03, 0.05, 0.065), PCT,
         "Management: mid single-digit growth a year to 2029"),
        ("af_m29", "Airframes margin reached by 2029", (0.085, 0.11, 0.125),
         PCT, "Management: low teens. 2025 actual 8.0%; 2026 guidance 8.6%"),
        ("vc29", "Non-cash VC as % of Engines revenue, 2029",
         (0.16, 0.11, 0.09), PCT,
         "2026 guidance: 20.6% (£360m of £1,750m). Lower = profit turning "
         "into cash faster"),
        ("vc35", "Non-cash VC as % of Engines revenue, 2035",
         (0.08, 0.05, 0.025), PCT,
         "Keeps falling as aftermarket cash from older engines catches up"),
        ("gg", "Remaining Garden Grove cash cost (£m)", (250, 120, 60), GBP,
         "Base: H2 exceptional costs £25-30m + compensation fund up to $100m "
         "(~£75m) + ~£15m legal. Bull: insurance pays half. Bear: lawsuits "
         "double it"),
        ("p", "Probability", (0.25, 0.50, 0.25), PCT,
         "Base case twice as likely as either tail"),
    ]
    r = 4
    for key, label, vals, fmt, why in levers:
        put(ws, f"A{r}", label)
        for col, v in zip("BCD", vals):
            put(ws, f"{col}{r}", v, fmt, BLUE, fill=YELLOW)
        put(ws, f"E{r}", why, color=GREY)
        SCEN_ROWS[key] = r
        r += 1
    put(ws, f"A{r}", "Probabilities sum to", italic=True)
    put(ws, f"B{r}", f"=SUM(B{SCEN_ROWS['p']}:D{SCEN_ROWS['p']})", PCT)
    r += 2
    section(ws, r, "Results", "E")
    r += 1
    out_rows = [
        ("value", "Value per share", PENCE, "{s}!$D$62"),
        ("upside", "Upside / (downside) vs share price", PCT, "{s}!$D$63"),
        ("fcf29", "2029 FCF after interest and tax (£m)", GBP, "{s}!$G$67"),
        ("fcf29_pct", "% of management's 2029 FCF target", PCT, "{s}!$G$69"),
        ("ev", "Enterprise value (£m)", GBP, "{s}!$D$56"),
    ]
    for key, label, fmt, ref in out_rows:
        put(ws, f"A{r}", label)
        for name, col in SCEN_COL.items():
            put(ws, f"{col}{r}", "=" + ref.format(s=f"DCF_{name}"), fmt, GREEN)
        SCEN_ROWS[key] = r
        r += 1
    r += 1
    put(ws, f"A{r}", "Probability-weighted value per share", bold=True)
    put(ws, f"B{r}",
        f"=SUMPRODUCT(B{SCEN_ROWS['p']}:D{SCEN_ROWS['p']},"
        f"B{SCEN_ROWS['value']}:D{SCEN_ROWS['value']})", PENCE, bold=True)
    SCEN_ROWS["weighted"] = r
    r += 1
    put(ws, f"A{r}", "Upside / (downside) vs share price", bold=True)
    put(ws, f"B{r}", f"=B{r - 1}/{I('price')}-1", PCT, bold=True)
    SCEN_ROWS["weighted_up"] = r
    widths(ws, {"A": 44, "B": 11, "C": 11, "D": 11, "E": 100})
    ws.freeze_panes = "B4"
    return ws


def S(key, scen):
    return f"Scenarios!${SCEN_COL[scen]}${SCEN_ROWS[key]}"


# ------------------------------------------------------------------ DCF
def build_dcf(wb, scen):
    ws = wb.create_sheet(f"DCF_{scen}")
    title(ws, f"Discounted cash flow: {scen.lower()} case",
          "£m unless stated. Forecasts at US$1.37/£, the 2026 guidance rate. "
          "Valued at 30 Jun 2026, the date of the latest net debt.")
    header_row(ws, 4, ["Year", "Note"] +
               [f"{y}{'A' if y == 2025 else 'E'}" for y in YEARS])
    first, last = COL[2026], COL[2035]

    def row(r, label, note="", fmt=GBP, bold=False):
        put(ws, f"A{r}", label, bold=bold)
        put(ws, f"B{r}", note, color=GREY)
        for c in COLS:
            ws[f"{c}{r}"].number_format = fmt
            ws[f"{c}{r}"].font = font(bold=bold)

    def link(ref, value, fmt):  # a cross-sheet link, shown green
        cell = ws[ref]
        cell.value = value
        cell.font = font(GREEN)
        cell.number_format = fmt

    def typed(ref, value, fmt):  # a historical figure typed from filings
        cell = ws[ref]
        cell.value = value
        cell.font = font(BLUE)
        cell.number_format = fmt

    # timing rows
    row(5, "Year (number)", fmt="0")
    row(6, "Years after 2029", fmt="0")
    row(7, "Progress to 2029 levels", "0 in 2026, 1 by 2029", fmt=PCT)
    row(8, "Discount period (years)", "From 30 Jun 2026, mid-period",
        fmt="0.00")
    for y, c in COL.items():
        typed(f"{c}5", y, "0")
        if y >= 2026:
            ws[f"{c}6"] = f"=MAX(0,{c}5-2029)"
            ws[f"{c}7"] = f"=MIN(MAX({c}5-2026,0),3)/3"
            ws[f"{c}8"] = 0.25 if y == 2026 else f"={c}5-2026"
            if y == 2026:
                ws[f"{c}8"].font = font(BLUE)

    # revenue
    row(10, "Engines revenue growth", "2027-29 lever, then fades", fmt=PCT)
    row(11, "Engines revenue", "2026 = guidance mid-point", bold=True)
    row(12, "Airframes revenue growth", "2027-29 lever, then fades", fmt=PCT)
    row(13, "Airframes revenue (before Garden Grove)", "2026 = guidance")
    row(14, "Garden Grove lost revenue", "2026 only")
    row(15, "Group revenue", bold=True)
    row(16, "Revenue growth", fmt=PCT)
    typed(f"{COL[2025]}11", 1632, GBP)
    typed(f"{COL[2025]}13", 1957, GBP)
    link(f"{COL[2026]}11", f"={I('g_eng_rev')}", GBP)
    link(f"{COL[2026]}13", f"={I('g_af_rev')}", GBP)
    ws[f"{COL[2026]}10"] = f"={COL[2026]}11/{COL[2025]}11-1"
    ws[f"{COL[2026]}12"] = f"={COL[2026]}13/{COL[2025]}13-1"
    link(f"{COL[2026]}14", f"=-{I('gg_rev')}", GBP)
    for y in YEARS[2:]:
        c, p = COL[y], COL[y - 1]
        for gr, lever in ((10, "eng_g"), (12, "af_g")):
            lev = S(lever, scen)
            ws[f"{c}{gr}"] = (f"={lev}" if y <= 2029 else
                              f"={lev}+({I('fade_to')}-{lev})*{c}6/6")
        ws[f"{c}11"] = f"={p}11*(1+{c}10)"
        ws[f"{c}13"] = f"={p}13*(1+{c}12)"
        ws[f"{c}14"] = 0
    for y in YEARS:
        c = COL[y]
        if y == 2025:
            ws[f"{c}14"] = 0
        ws[f"{c}15"] = f"={c}11+{c}13+{c}14"
        if y > 2025:
            ws[f"{c}16"] = f"={c}15/{COL[y - 1]}15-1"

    # profit
    row(18, "Engines margin", "Straight line to the 2029 lever", fmt=PCT)
    row(19, "Engines adjusted operating profit")
    row(20, "Airframes margin", "Straight line to the 2029 lever", fmt=PCT)
    row(21, "Airframes adjusted operating profit")
    row(22, "Garden Grove lost profit", "2026 only")
    row(23, "Central costs", "Grow with inflation")
    row(24, "Group adjusted operating profit", bold=True)
    row(25, "Adjusted operating margin", fmt=PCT)
    c25, c26 = COL[2025], COL[2026]
    typed(f"{c25}19", 520, GBP)
    typed(f"{c25}21", 156, GBP)
    typed(f"{c25}23", -29, GBP)
    ws[f"{c25}18"] = f"={c25}19/{c25}11"
    ws[f"{c25}20"] = f"={c25}21/{c25}13"
    ws[f"{c25}22"] = 0
    ws[f"{c26}18"] = f"={I('g_eng_op')}/{I('g_eng_rev')}"
    ws[f"{c26}20"] = f"={I('g_af_op')}/{I('g_af_rev')}"
    link(f"{c26}22", f"=-{I('gg_op')}", GBP)
    link(f"{c26}23", f"=-{I('g_corp')}", GBP)
    for y in YEARS[1:]:
        c = COL[y]
        if y > 2026:
            ws[f"{c}18"] = f"=${c26}18+({S('eng_m29', scen)}-${c26}18)*{c}7"
            ws[f"{c}20"] = f"=${c26}20+({S('af_m29', scen)}-${c26}20)*{c}7"
            ws[f"{c}22"] = 0
            ws[f"{c}23"] = f"={COL[y - 1]}23*(1+{I('corp_g')})"
        ws[f"{c}19"] = f"={c}11*{c}18"
        ws[f"{c}21"] = f"={c}13*{c}20"
    for c in COLS:
        ws[f"{c}24"] = f"={c}19+{c}21+{c}22+{c}23"
        ws[f"{c}25"] = f"={c}24/{c}15"

    # cash flow
    row(27, "Depreciation and amortisation")
    row(28, "Adjusted EBITDA", bold=True)
    row(29, "Lease payments")
    row(30, "Capital expenditure")
    row(31, "Working capital (excluding VC)", "2025: stock and debtors")
    row(32, "Non-cash VC as % of Engines revenue", "Lever for 2029 and 2035",
        fmt=PCT)
    row(33, "Non-cash variable consideration", "Profit not yet received")
    row(34, "GTF powder-metal payments")
    row(35, "Pension contributions")
    row(36, "Restructuring and other", "2025: incl. contract releases")
    row(37, "Cash flow before tax", "2025 ties to £252m reported", bold=True)
    row(38, "Tax", "Rate x (profit less non-cash VC)")
    row(39, "Unlevered free cash flow (FCFF)", bold=True)
    row(40, "Less: H1 2026 (already in net debt)")
    row(41, "Cash flow to discount", bold=True)
    row(42, "Discount factor", fmt=FACTOR)
    row(43, "Present value")
    for ref, v in (("27", 138), ("29", -31), ("30", -94), ("31", 33),
                   ("33", -324), ("34", -68), ("35", -22), ("36", -27)):
        typed(f"{c25}{ref}", v, GBP)
    ws[f"{c25}32"] = f"=-{c25}33/{c25}11"
    ws[f"{c26}32"] = f"={I('g_vc')}/{I('g_eng_rev')}"
    sched = {2026: ("gtf26", "ro26"), 2027: ("gtf27", "ro27"),
             2028: ("gtf28", "ro28")}
    for y in YEARS[1:]:
        c, p = COL[y], COL[y - 1]
        ws[f"{c}27"] = f"={I('da_pct')}*{c}15"
        ws[f"{c}29"] = f"=-{I('lease_pct')}*{c}15"
        ws[f"{c}30"] = f"=-{I('capex_pct')}*{c}15"
        ws[f"{c}31"] = f"=-{I('wc_pct')}*({c}15-{p}15)"
        if 2027 <= y <= 2029:
            ws[f"{c}32"] = (f"=${c26}32+({S('vc29', scen)}-${c26}32)"
                            f"*{c}7")
        elif y > 2029:
            g29 = f"${COL[2029]}$32"
            ws[f"{c}32"] = f"={g29}+({S('vc35', scen)}-{g29})*{c}6/6"
        ws[f"{c}33"] = f"=-{c}32*{c}11"
        if y in sched:
            link(f"{c}34", f"=-{I(sched[y][0])}", GBP)
            link(f"{c}36", f"=-{I(sched[y][1])}", GBP)
        else:
            ws[f"{c}34"] = 0
            ws[f"{c}36"] = 0
        ws[f"{c}35"] = f"=-{I('pen_cash')}"
        ws[f"{c}38"] = f"=-{I('tax')}*MAX(0,{c}24+{c}33)"
        ws[f"{c}39"] = f"={c}37+{c}38"
        ws[f"{c}40"] = f"=-{I('h1_fcff')}" if y == 2026 else 0
        ws[f"{c}41"] = f"={c}39+{c}40"
        ws[f"{c}42"] = f"=1/(1+{I('wacc')})^{c}8"
        ws[f"{c}43"] = f"={c}41*{c}42"
    for c in COLS:
        ws[f"{c}28"] = f"={c}24+{c}27"
        ws[f"{c}37"] = (f"={c}24+{c}27+{c}29+{c}30+{c}31+{c}33+{c}34"
                        f"+{c}35+{c}36")

    # valuation block, values in column D
    section(ws, 45, "Valuation", "M")
    val = [
        (46, "WACC", f"={I('wacc')}", PCT2),
        (47, "Terminal growth", f"={I('g')}", PCT2),
        (48, "FCFF in 2036", f"={last}39*(1+D47)", GBP),
        (49, "Terminal value at end-2035", "=D48/(D46-D47)", GBP),
        (50, "Years to end-2035", f"={last}8+0.5", "0.00"),
        (51, "Discount factor", "=1/(1+D46)^D50", FACTOR),
        (52, "Present value of terminal value", "=D49*D51", GBP),
        (53, "Present value of 2026-35 cash flows",
         f"=SUM({first}43:{last}43)", GBP),
        (56, "Enterprise value", "=D52+D53", GBP),
        (57, "Less: net debt, 30 Jun 2026", f"=-{I('net_debt')}", GBP),
        (58, "Less: pension deficit", f"=-{I('pension')}", GBP),
        (59, "Less: remaining Garden Grove costs", f"=-{S('gg', scen)}", GBP),
        (60, "Equity value", "=SUM(D56:D59)", GBP),
        (61, "Shares in issue (m)", f"={I('shares')}", GBP),
        (62, "Value per share", "=D60/D61*100", PENCE),
        (63, "Upside / (downside) vs share price", f"=D62/{I('price')}-1",
         PCT),
        (64, "Terminal value as % of enterprise value", "=D52/D56", PCT),
        (65, "Implied exit multiple (EV / 2035 EBITDA)", f"=D49/{last}28",
         MULT),
    ]
    for r, label, f, fmt in val:
        bold = r in (56, 60, 62)
        put(ws, f"A{r}", label, bold=bold)
        cell = put(ws, f"D{r}", f, fmt, bold=bold)
        if "Inputs!" in f or "Scenarios!" in f:
            cell.font = font(GREEN, bold)
    put(ws, "B48", "2035 FCFF grown once", color=GREY)
    put(ws, "B50", "Mid-2026 to end-2035", color=GREY)
    put(ws, "B59", "Scenario lever", color=GREY)

    # memo: compare with management's free cash flow target
    section(ws, 66, "Memo: cash flow on Melrose's own definition", "M")
    row(67, "FCF after interest and tax", "FCFF less after-tax interest")
    row(68, "Management 2029 target at US$1.37")
    row(69, "Model as % of target", fmt=PCT)
    for y in YEARS[1:]:
        c = COL[y]
        ws[f"{c}67"] = f"={c}39-{I('kd')}*{I('net_debt')}"
    link(f"{COL[2029]}68", f"={I('t_fcf_137')}", GBP)
    ws[f"{COL[2029]}69"] = f"={COL[2029]}67/{COL[2029]}68"

    for c in COLS:
        for r in (11, 15, 24, 28, 37, 39, 41):
            ws[f"{c}{r}"].border = Border(top=THIN)
    widths(ws, {"A": 40, "B": 30, **{c: 9.5 for c in COLS}})
    ws.freeze_panes = "C5"
    return ws


# ---------------------------------------------------------- Sensitivity
def build_sensitivity(wb):
    ws = wb.create_sheet("Sensitivity")
    title(ws, "Base case: value per share (p) by WACC and terminal growth",
          "Each cell re-discounts the base-case cash flows. Rows move WACC "
          "around the model's rate; columns move terminal growth.")
    put(ws, "A4", "WACC \\ terminal growth", bold=True)
    gs = [0.015, 0.02, 0.025, 0.03, 0.035]
    steps = [-0.01, -0.005, 0, 0.005, 0.01]
    for j, g in enumerate(gs):
        put(ws, f"{get_column_letter(2 + j)}4", g, PCT, BLUE, bold=True,
            align="center")
    b = "DCF_Base"
    for i, step in enumerate(steps):
        r = 5 + i
        put(ws, f"A{r}", f"={I('wacc')}+({step})", PCT2, GREEN, bold=True)
        for j in range(len(gs)):
            col = get_column_letter(2 + j)
            w, g = f"$A{r}", f"{col}$4"
            f = (f"=(SUMPRODUCT({b}!$D$41:$M$41/(1+{w})^{b}!$D$8:$M$8)"
                 f"+{b}!$M$39*(1+{g})/({w}-{g})/(1+{w})^{b}!$D$50"
                 f"-{I('net_debt')}-{I('pension')}-{S('gg', 'Base')})"
                 f"/{I('shares')}*100")
            cell = put(ws, f"{col}{r}", f, PENCE)
            if step == 0 and g == 0.025:
                cell.fill = YELLOW
                cell.font = font(bold=True)
    put(ws, "A11", "Share price", italic=True)
    put(ws, "B11", f"={I('price')}", PENCE1, GREEN)
    put(ws, "A12", "Cells above the share price", italic=True)
    put(ws, "B12", '=COUNTIF(B5:F9,">"&B11)&" of 25"')
    widths(ws, {"A": 26, **{get_column_letter(2 + j): 11 for j in range(5)}})
    return ws


# ---------------------------------------------------------------- Comps
def build_comps(wb):
    ws = wb.create_sheet("Comps")
    title(ws, "Cross-checks: peer multiples and a recent takeover",
          "Consensus adjusted EPS from S&P Global via stockanalysis.com. "
          "Prices are the latest on the page as of 6-8 Oct 2026.")
    header_row(ws, 4, ["Company", "Ticker", "Currency", "Share price",
                       "Price date", "EPS 2026E", "P/E 2026E", "Notes"])
    peers = [
        ("Rolls-Royce", "RR.L", "pence", 1369.8, "8 Oct 2026", 44.0,
         "Engines; services earned per flying hour"),
        ("Safran", "SAF.PA", "EUR", 328.90, "6 Oct 2026", 10.42,
         "Engines (CFM with GE) and equipment"),
        ("MTU Aero Engines", "MTX.DE", "EUR", 367.10, "6 Oct 2026", 18.98,
         "Engine parts partner, like GKN's RRSP model"),
        ("GE Aerospace", "GE", "USD", 303.62, "7 Oct 2026", 7.92,
         "Engines (LEAP, GEnx, GE9X)"),
        ("Howmet Aerospace", "HWM", "USD", 231.60, "6 Oct 2026", 5.34,
         "Engine and airframe components"),
    ]
    r = 5
    for name, tick, ccy, px, date, eps, note in peers:
        for col, v in zip("ABC", (name, tick, ccy)):
            put(ws, f"{col}{r}", v)
        put(ws, f"D{r}", px, "#,##0.00", BLUE)
        put(ws, f"E{r}", date, color=BLUE)
        put(ws, f"F{r}", eps, "#,##0.00", BLUE)
        put(ws, f"G{r}", f"=D{r}/F{r}", MULT)
        put(ws, f"H{r}", note, color=GREY)
        r += 1
    put(ws, f"A{r}", "Peer median", bold=True)
    put(ws, f"G{r}", f"=MEDIAN(G5:G{r - 1})", MULT, bold=True)
    med = r
    r += 1
    put(ws, f"A{r}", "Melrose", bold=True)
    put(ws, f"B{r}", "MRO.L")
    put(ws, f"C{r}", "pence")
    put(ws, f"D{r}", f"={I('price')}", "#,##0.00", GREEN)
    put(ws, f"E{r}", "7 Oct 2026")
    put(ws, f"F{r}", f"={I('eps26')}", "#,##0.00", GREEN)
    put(ws, f"G{r}", f"=D{r}/F{r}", MULT, bold=True)
    mro = r
    r += 1
    put(ws, f"A{r}", "Melrose discount to peer median")
    put(ws, f"G{r}", f"=1-G{mro}/G{med}", PCT)
    r += 1
    put(ws, f"A{r}", "Melrose EPS excluding non-cash VC")
    put(ws, f"F{r}",
        f"=F{mro}-{I('g_vc')}*(1-{I('tax')})/{I('shares')}*100",
        "#,##0.00")
    put(ws, f"G{r}", f"=D{mro}/F{r}", MULT, bold=True)
    put(ws, f"H{r}", "On cash-backed earnings Melrose is not cheap: the "
        "debate is when VC turns into cash", color=GREY)
    cash_pe = r

    r += 2
    section(ws, r, "Takeover cross-check: Senior plc (UK aerospace and "
            "defence parts maker)", "H")
    r += 1
    rows = [
        ("senior", "Senior plc: EV / FY2025 adjusted EBITDA agreed", 15.2, MULT,
         True, "Blackstone/Tinicum offer, Rule 2.7 announcement, 7 Apr 2026 "
         "(300p a share, EV ~£1,399m)"),
        ("ebitda25", "Melrose FY2025 adjusted EBITDA", "=DCF_Base!C28", GBP,
         False, "Adjusted operating profit £647m + D&A £138m"),
        ("vc25", "of which non-cash VC", "=-DCF_Base!C33", GBP, False,
         "FY2025 results"),
        ("cash25", "Melrose FY2025 EBITDA excluding VC", None, GBP, False,
         "Formula"),
        ("mult_rep", "Melrose EV / EBITDA", None, MULT, False,
         "EV includes leases, matching EBITDA"),
        ("mult_cash", "Melrose EV / EBITDA excluding VC", None, MULT, False,
         "Close to the multiple agreed for Senior"),
        ("v_cash", "Value at Senior's multiple, VC worth nothing", None,
         PENCE, False, "Applies 15.2x to cash EBITDA only"),
        ("v_full", "Value at Senior's multiple, VC worth full value", None,
         PENCE, False, "Applies 15.2x to all EBITDA"),
    ]
    at = {}
    for key, label, v, fmt, typed, note in rows:
        at[key] = r
        r += 1
    claims = (f"{I('net_debt')}+{I('leases')}+{I('pension')}"
              f"+{S('gg', 'Base')}")
    formulas = {
        "cash25": f"=G{at['ebitda25']}-G{at['vc25']}",
        "mult_rep": f"={I('ev')}/G{at['ebitda25']}",
        "mult_cash": f"={I('ev')}/G{at['cash25']}",
        "v_cash": (f"=(G{at['senior']}*G{at['cash25']}-({claims}))"
                   f"/{I('shares')}*100"),
        "v_full": (f"=(G{at['senior']}*G{at['ebitda25']}-({claims}))"
                   f"/{I('shares')}*100"),
    }
    for key, label, v, fmt, typed, note in rows:
        rr = at[key]
        put(ws, f"A{rr}", label, bold=key in ("v_cash", "v_full"))
        value = formulas.get(key, v)
        color = BLUE if typed else (GREEN if "!" in str(value) else BLACK)
        put(ws, f"G{rr}", value, fmt, color, bold=key in ("v_cash", "v_full"))
        put(ws, f"H{rr}", note, color=GREY)
    widths(ws, {"A": 46, "B": 9, "C": 9, "D": 12, "E": 12, "F": 11,
                "G": 11, "H": 70})
    return ws, at, cash_pe, med, mro


# -------------------------------------------------------------- History
def build_history(wb):
    ws = wb.create_sheet("History")
    title(ws, "Reported figures",
          "FY2024-25 from the FY2025 results (27 Feb 2026); half years from "
          "the H1 2026 results (31 Jul 2026). £m unless stated.")
    header_row(ws, 4, ["Item", "FY2024", "FY2025", "H1 2025", "H1 2026"])
    data = [
        ("Engines revenue", (1459, 1632, 781, 896)),
        ("Airframes revenue", (2009, 1957, 939, 977)),
        ("Group revenue", "sum2"),
        ("Engines adjusted operating profit", (422, 520, 261, 303)),
        ("Airframes adjusted operating profit", (144, 156, 63, 62)),
        ("Central costs", (-26, -29, -14, -18)),
        ("Group adjusted operating profit", "sum3"),
        ("Engines margin", "m_eng"),
        ("Airframes margin", "m_af"),
        ("Group margin", "m_grp"),
        ("Variable consideration in profit (non-cash)", (274, 324, 182, 206)),
        ("VC as % of group profit", "vc_share"),
        ("Free cash flow (after interest and tax)", (-74, 125, -54, 13)),
        ("Net debt (excluding leases)", (1321, 1407, None, 1530)),
        ("Shares in issue excluding treasury (m)", (1286, 1258, None, 1246)),
        ("Adjusted diluted EPS (p)", (26.4, 32.1, 15.1, 17.7)),
        ("Dividend per share (p)", (6.0, 7.2, 2.4, 2.7)),
    ]
    r = 5
    rows = {}
    for label, vals in data:
        put(ws, f"A{r}", label)
        rows[label] = r
        for j, col in enumerate("BCDE"):
            if isinstance(vals, tuple):
                if vals[j] is not None:
                    fmt = "0.0" if "(p)" in label else GBP
                    put(ws, f"{col}{r}", vals[j], fmt, BLUE)
            elif vals == "sum2":
                put(ws, f"{col}{r}", f"={col}5+{col}6", GBP, bold=True)
            elif vals == "sum3":
                put(ws, f"{col}{r}", f"={col}8+{col}9+{col}10", GBP,
                    bold=True)
            elif vals == "m_eng":
                put(ws, f"{col}{r}", f"={col}8/{col}5", PCT)
            elif vals == "m_af":
                put(ws, f"{col}{r}", f"={col}9/{col}6", PCT)
            elif vals == "m_grp":
                put(ws, f"{col}{r}", f"={col}11/{col}7", PCT)
            elif vals == "vc_share":
                put(ws, f"{col}{r}", f"={col}15/{col}11", PCT)
        r += 1
    put(ws, f"A{r + 1}", "Central costs are group profit less the two "
        "divisions. H1 dividend is the interim.", color=GREY, italic=True)
    widths(ws, {"A": 46, "B": 11, "C": 11, "D": 11, "E": 11})
    return ws


# -------------------------------------------------------------- Summary
def build_summary(wb, comps_at, cash_pe, med, mro):
    ws = wb["Summary"]
    title(ws, "Melrose Industries (MRO.L): equity research model",
          "Rohan Sharma, 8 Oct 2026. Student project for learning; not "
          "investment advice.")
    w = SCEN_ROWS["weighted"]
    put(ws, "A4", "Recommendation", bold=True)
    put(ws, "B4", f'=IF(Scenarios!B{w + 1}>=0.15,"BUY",'
        f'IF(Scenarios!B{w + 1}<=-0.1,"SELL","HOLD"))', bold=True)
    put(ws, "C4", "Buy if fair value is 15%+ above the price; sell if 10%+ "
        "below", color=GREY)
    lines = [
        ("Share price (7 Oct 2026)", f"={I('price')}", PENCE1),
        ("Fair value (probability-weighted)", f"=Scenarios!B{w}", PENCE),
        ("Upside", f"=Scenarios!B{w + 1}", PCT),
        ("Market capitalisation (£m)", f"={I('mcap')}", GBP),
        ("Enterprise value incl. leases (£m)", f"={I('ev')}", GBP),
        ("P/E 2026E (consensus EPS)", f"=Comps!G{mro}", MULT),
        ("Peer median P/E 2026E", f"=Comps!G{med}", MULT),
        ("P/E 2026E on EPS excluding non-cash VC", f"=Comps!G{cash_pe}",
         MULT),
        ("FCF yield, 2026 guidance mid-point", f"={I('g_fcf')}/{I('mcap')}",
         PCT),
        ("FCF yield on the 2029 target (at US$1.37)",
         f"={I('t_fcf_137')}/{I('mcap')}", PCT),
        ("WACC", f"={I('wacc')}", PCT2),
    ]
    r = 6
    for label, f, fmt in lines:
        put(ws, f"A{r}", label)
        put(ws, f"B{r}", f, fmt, GREEN)
        r += 1
    r += 1
    header_row(ws, r, ["Scenario", "Value", "Upside", "Probability",
                       "2029 FCF (£m)", "% of 2029 target"])
    r += 1
    for name, col in SCEN_COL.items():
        put(ws, f"A{r}", name)
        for c2, key, fmt in (("B", "value", PENCE), ("C", "upside", PCT),
                             ("D", "p", PCT), ("E", "fcf29", GBP),
                             ("F", "fcf29_pct", PCT)):
            put(ws, f"{c2}{r}", f"=Scenarios!{col}{SCEN_ROWS[key]}", fmt,
                GREEN)
        r += 1
    r += 1
    header_row(ws, r, ["Valuation range (football field)", "Low (p)",
                       "High (p)", "Basis"])
    r += 1
    sens = "Sensitivity"
    ff = [
        ("52-week trading range", f"={I('low52')}", f"={I('high52')}",
         "lse.co.uk"),
        ("Analyst price targets", f"={I('tgt_low')}", f"={I('tgt_high')}",
         "15 analysts, 7 Oct 2026"),
        ("Senior plc takeover multiple (15.2x EBITDA)",
         f"=Comps!G{comps_at['v_cash']}", f"=Comps!G{comps_at['v_full']}",
         "Low: VC worth nothing. High: VC worth full value"),
        ("DCF, base case (WACC +/-0.5pt, growth +/-0.5pt)",
         f"=MIN({sens}!C6:E8)", f"=MAX({sens}!C6:E8)", "Sensitivity sheet"),
        ("DCF, bear to bull", f"=Scenarios!B{SCEN_ROWS['value']}",
         f"=Scenarios!D{SCEN_ROWS['value']}", "Scenarios sheet"),
    ]
    for label, lo, hi, basis in ff:
        put(ws, f"A{r}", label)
        put(ws, f"B{r}", lo, PENCE, GREEN)
        put(ws, f"C{r}", hi, PENCE, GREEN)
        put(ws, f"D{r}", basis, color=GREY)
        r += 1
    r += 1
    notes = [
        "How to use: change the yellow levers on Scenarios or any blue input "
        "on Inputs; every value recalculates.",
        "Colour code: blue = typed-in data, black = formula, green = link to "
        "another sheet, yellow = key judgement.",
        "The key judgement is how fast non-cash variable consideration (VC) "
        "turns into cash: see DCF row 32.",
    ]
    for n in notes:
        put(ws, f"A{r}", n, color=GREY, italic=True)
        r += 1
    widths(ws, {"A": 46, "B": 12, "C": 12, "D": 12, "E": 14, "F": 16})
    return ws


def main():
    wb = Workbook()
    wb.active.title = "Summary"
    build_inputs(wb)
    build_scenarios(wb)
    for scen in ("Base", "Bear", "Bull"):
        build_dcf(wb, scen)
    build_sensitivity(wb)
    _, comps_at, cash_pe, med, mro = build_comps(wb)
    build_history(wb)
    build_summary(wb, comps_at, cash_pe, med, mro)
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
