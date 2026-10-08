"""Check the Excel model against the independent Python replica.

The workbook stores LibreOffice/Excel-calculated values, so this reads
them back and compares every forecast line, valuation output and
sensitivity cell with python/dcf.py.

Run:  python3 -m pytest python/test_model.py   (or python3 python/test_model.py)
"""
from pathlib import Path

from openpyxl import load_workbook

import dcf

BOOK = Path(__file__).resolve().parents[1] / "model" / "Melrose_DCF.xlsx"
COLS = dict(zip(range(2025, 2036), "CDEFGHIJKLM"))
ROWS = {15: "rev", 24: "op", 28: "ebitda", 33: "vc", 37: "pre_tax",
        38: "tax", 39: "fcff", 41: "cash_flow", 42: "df", 43: "pv",
        67: "fcf_levered"}
VALUATION = {49: "tv", 52: "pv_tv", 56: "ev", 60: "equity", 62: "value_p",
             64: "tv_share", 65: "exit_multiple"}


def close(a, b, tol=1e-6):
    return a is not None and abs(a - b) <= tol * max(1.0, abs(b))


def test_workbook_matches_python():
    wb = load_workbook(BOOK, data_only=True)
    checked = 0
    for scen, params in dcf.SCENARIOS.items():
        ws, r = wb[f"DCF_{scen}"], dcf.run(params)
        for row, key in ROWS.items():
            for y in range(2026, 2036):
                assert close(ws[f"{COLS[y]}{row}"].value, r[key][y]), \
                    (scen, key, y)
                checked += 1
        for row, key in VALUATION.items():
            assert close(ws[f"D{row}"].value, r[key]), (scen, key)
            checked += 1
    sens = wb["Sensitivity"]
    for i in range(5):
        rate = sens[f"A{5 + i}"].value
        for col in "BCDEF":
            g = sens[f"{col}4"].value
            want = dcf.run(dcf.SCENARIOS["Base"], rate, g)["value_p"]
            assert close(sens[f"{col}{5 + i}"].value, want), (rate, g)
            checked += 1
    return checked


def test_2025_ties_to_reported_cash_flow():
    ws = load_workbook(BOOK, data_only=True)["DCF_Base"]
    assert ws["C24"].value == 647   # adjusted operating profit, FY2025
    assert ws["C37"].value == 252   # free cash flow before interest and tax


if __name__ == "__main__":
    n = test_workbook_matches_python()
    test_2025_ties_to_reported_cash_flow()
    print(f"OK: {n} workbook values match the Python model")
